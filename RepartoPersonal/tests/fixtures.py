"""Utilidades de prueba: configuración FICTICIA (nunca datos reales)."""

from pathlib import Path

import openpyxl

from reparto.datos import POSITIONS


def pattern_sequence(kind: str) -> list[str]:
    """Secuencia de 168 días (24 semanas): L-V con turno fijo, S-D descanso.

    kind: 'M' mañana siempre, 'T' tarde siempre, 'A' alterna semanas M/T.
    """
    marks = []
    for day in range(168):
        weekday = day % 7
        week = day // 7
        if weekday >= 5:
            marks.append("D")
        elif kind == "A":
            marks.append("M" if week % 2 == 0 else "T")
        else:
            marks.append(kind)
    return marks


def write_config(directory: Path, workers: list[dict]) -> tuple[Path, Path]:
    """workers: [{'name', 'seq': list[168], 'skills': iterable, 'ranking': {pos: n}, 'baja': bool}]."""
    directory.mkdir(parents=True, exist_ok=True)

    wb_seq = openpyxl.Workbook()
    ws = wb_seq.active
    ws.title = "Secuencias de Turnos"
    ws.cell(1, 1, "SECUENCIA BASE DE TURNOS")
    ws.cell(3, 1, "N.º")
    ws.cell(3, 2, "Apellidos y Nombre")
    for i, w in enumerate(workers):
        r = 4 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, w["name"])
        for c, mark in enumerate(w["seq"]):
            ws.cell(r, 3 + c, mark)
    seq_path = directory / "secuencias.xlsx"
    wb_seq.save(seq_path)

    wb_pref = openpyxl.Workbook()
    wp = wb_pref.active
    wp.title = "Trabajadores y Preferencias"
    wp.cell(3, 1, "N.º")
    wp.cell(3, 2, "Apellidos y Nombre")
    wp.cell(3, 3, "BAJA")
    for idx, pos in enumerate(POSITIONS):
        wp.cell(3, 4 + idx, pos)
        wp.cell(3, 20 + idx, pos)
    for i, w in enumerate(workers):
        r = 4 + i
        wp.cell(r, 1, i + 1)
        wp.cell(r, 2, w["name"])
        if w.get("baja"):
            wp.cell(r, 3, "X")
        for idx, pos in enumerate(POSITIONS):
            if pos in w.get("skills", ()):
                wp.cell(r, 4 + idx, "X")
            if pos in w.get("ranking", {}):
                wp.cell(r, 20 + idx, w["ranking"][pos])
    pref_path = directory / "preferencias.xlsx"
    wb_pref.save(pref_path)
    return seq_path, pref_path
